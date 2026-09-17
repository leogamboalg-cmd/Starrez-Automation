"""Shared response loading and lightweight JSON processing state."""

from __future__ import annotations

import csv
import json
import os
from datetime import datetime, timezone
from pathlib import Path


APP_NAME = "Add New Contacts for Newsletter"
APP_DATA_ROOT = Path(
    os.environ.get("LOCALAPPDATA", Path.home() / "AppData" / "Local")
)
SETTINGS_FILE = APP_DATA_ROOT / APP_NAME / "launcher_settings.json"
ONEDRIVE_ROOT = Path(
    os.environ.get(
        "OneDriveCommercial",
        Path.home() / "OneDrive - Cal Poly Pomona",
    )
)
DEFAULT_SHARED_RESPONSE_FILE = ONEDRIVE_ROOT / "UHS Newsletter Reponses.csv"
STATE_FILE = APP_DATA_ROOT / APP_NAME / "response_state.json"

FIELD_ALIASES = {
    "id": ("id", "response id", "responseid"),
    "start_time": ("start time", "starttime"),
    "completion_time": ("completion time", "completiontime"),
    "first": ("contact first name", "first", "first name"),
    "last": ("contact last name", "last", "last name"),
    "email": ("contact email", "email", "email address"),
}


def get_shared_response_file():
    """Return the shared response file selected during installation."""
    try:
        settings = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
        configured_path = settings.get("shared_response_file")
        if isinstance(configured_path, str) and configured_path.strip():
            return Path(configured_path.strip())
    except (OSError, json.JSONDecodeError):
        pass
    return DEFAULT_SHARED_RESPONSE_FILE


# Kept for compatibility with older callers; new code should use the function so
# a settings change is picked up without relying on an import-time value.
SHARED_RESPONSE_FILE = get_shared_response_file()


def _normalized(value):
    return " ".join(str(value or "").strip().lower().split())


def _column_map(fieldnames):
    available = {_normalized(name): name for name in fieldnames if name}
    result = {}
    for key, aliases in FIELD_ALIASES.items():
        result[key] = next(
            (available[alias] for alias in aliases if alias in available),
            None,
        )
    missing = [key for key in ("first", "last", "email") if not result[key]]
    if missing:
        raise ValueError(
            "The contact file is missing required columns for: "
            + ", ".join(missing)
        )
    return result


def _contact_from_values(values, columns, row_number):
    response_id = str(values.get(columns["id"], "") or "").strip() if columns["id"] else ""
    if not response_id:
        response_id = f"row-{row_number}"
    return {
        "id": response_id,
        "row_number": row_number,
        "start_time": str(values.get(columns["start_time"], "") or "").strip()
        if columns["start_time"] else "",
        "completion_time": str(values.get(columns["completion_time"], "") or "").strip()
        if columns["completion_time"] else "",
        "first": str(values.get(columns["first"], "") or "").strip(),
        "last": str(values.get(columns["last"], "") or "").strip(),
        "email": str(values.get(columns["email"], "") or "").strip(),
    }


def load_csv_contacts(path):
    contacts = []
    with Path(path).open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        if not reader.fieldnames:
            raise ValueError("The CSV file does not contain a header row.")
        columns = _column_map(reader.fieldnames)
        for row_number, values in enumerate(reader, start=2):
            contact = _contact_from_values(values, columns, row_number)
            if contact["first"] and contact["last"] and contact["email"]:
                contacts.append(contact)
    return contacts


def load_excel_contacts(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        columns = _column_map(headers)
        header_positions = {
            key: headers.index(value) + 1 if value is not None else None
            for key, value in columns.items()
        }
        contacts = []
        for row_number in range(2, sheet.max_row + 1):
            values = {
                header: sheet.cell(row=row_number, column=position).value
                for header, position in zip(headers, range(1, len(headers) + 1))
                if header is not None
            }
            contact = _contact_from_values(values, columns, row_number)
            if contact["first"] and contact["last"] and contact["email"]:
                contacts.append(contact)
        return contacts
    finally:
        workbook.close()


def load_contacts(path):
    path = Path(path)
    if path.suffix.lower() == ".csv":
        return load_csv_contacts(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return load_excel_contacts(path)
    raise ValueError("Choose a CSV or Excel .xlsx file.")


def load_state():
    try:
        state = json.loads(STATE_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        state = {}
    processed = state.get("processed_ids", {})
    state["processed_ids"] = processed if isinstance(processed, dict) else {}
    notified = state.get("last_notified_ids", [])
    state["last_notified_ids"] = notified if isinstance(notified, list) else []
    return state


def save_state(state):
    STATE_FILE.parent.mkdir(parents=True, exist_ok=True)
    temporary = STATE_FILE.with_suffix(".tmp")
    temporary.write_text(json.dumps(state, indent=2), encoding="utf-8")
    temporary.replace(STATE_FILE)


def pending_contacts(contacts):
    processed = load_state()["processed_ids"]
    return [contact for contact in contacts if contact["id"] not in processed]


def mark_processed(contact):
    state = load_state()
    state["processed_ids"][contact["id"]] = {
        "processed_at": datetime.now(timezone.utc).isoformat(),
        "completion_time": contact["completion_time"],
    }
    save_state(state)


def notification_needed(pending):
    state = load_state()
    pending_ids = [contact["id"] for contact in pending]
    return bool(pending_ids) and state["last_notified_ids"] != pending_ids


def record_notification(pending):
    state = load_state()
    state["last_notified_ids"] = [contact["id"] for contact in pending]
    save_state(state)


def reset_notification_state():
    state = load_state()
    state["last_notified_ids"] = []
    save_state(state)
