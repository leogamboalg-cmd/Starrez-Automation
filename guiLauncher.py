from pathlib import Path
import json
import queue
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import pyautogui
from openpyxl import load_workbook
from PIL import Image, ImageTk

from guiFinished import get_column_numbers, process_contact
from starrez_screen import STARREZ_URL, check_starrez_ready


APP_BACKGROUND = "#F4F7FB"
CARD_BACKGROUND = "#FFFFFF"
FIELD_BACKGROUND = "#F8FAFC"
BORDER = "#DDE5EF"
ACCENT = "#2563EB"
ACCENT_HOVER = "#1D4ED8"
ACCENT_SOFT = "#EAF1FF"
PRIMARY_TEXT = "#172033"
SECONDARY_TEXT = "#667085"
WARNING = "#B45309"
WARNING_SOFT = "#FFF7E8"
ERROR = "#C2414F"
ERROR_SOFT = "#FFF0F2"
SUCCESS = "#18794E"
SUCCESS_SOFT = "#EAF8F1"
SETTINGS_FILE = Path(__file__).with_name("launcher_settings.json")


class AutomationLauncher(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("StarRez Contact Runner")
        self.geometry("920x680")
        self.minsize(520, 420)
        self.configure(bg=APP_BACKGROUND)

        self.excel_path = tk.StringVar(
            value=self._load_saved_excel_path()
        )
        self.current_contact = tk.StringVar(value="No contact selected")
        self.status_text = tk.StringVar(value="Ready")
        self.progress_text = tk.StringVar(value="0 / 0")
        self.event_queue = queue.Queue()
        self.running = False
        self.website_retry = threading.Event()
        self.cancel_run = threading.Event()
        self.website_prompt = None
        self.protocol("WM_DELETE_WINDOW", self._close_app)

        self._configure_styles()
        self._build_interface()
        self.bind("<Configure>", self._on_resize)
        self.after(100, self._process_events)

    @staticmethod
    def _load_saved_excel_path():
        default_path = str(Path(__file__).with_name("contacts.xlsx"))
        try:
            settings = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
            saved_path = settings.get("excel_file")
            if isinstance(saved_path, str) and saved_path.strip():
                return saved_path
        except (OSError, json.JSONDecodeError):
            pass
        return default_path

    def _save_excel_path(self):
        settings = {"excel_file": self.excel_path.get().strip()}
        try:
            SETTINGS_FILE.write_text(
                json.dumps(settings, indent=2),
                encoding="utf-8",
            )
        except OSError as error:
            messagebox.showwarning(
                "Could not save preference",
                f"The workbook is selected, but its location could not be saved:\n{error}",
            )

    def _configure_styles(self):
        style = ttk.Style(self)
        style.theme_use("clam")
        style.configure(
            "Runner.Horizontal.TProgressbar",
            troughcolor="#E8EDF4",
            background=ACCENT,
            bordercolor="#E8EDF4",
            lightcolor=ACCENT,
            darkcolor=ACCENT,
            thickness=10,
        )

    def _build_interface(self):
        page = tk.Frame(self, bg=APP_BACKGROUND)
        self.main_page = page
        self.example_page = None
        page.pack(fill="both", expand=True)
        page.grid_rowconfigure(0, weight=1)
        page.grid_columnconfigure(0, weight=1)

        self.canvas = tk.Canvas(
            page,
            bg=APP_BACKGROUND,
            highlightthickness=0,
            bd=0,
        )
        scrollbar = ttk.Scrollbar(
            page,
            orient="vertical",
            command=self.canvas.yview,
        )
        self.canvas.configure(yscrollcommand=scrollbar.set)
        self.canvas.grid(row=0, column=0, sticky="nsew")
        scrollbar.grid(row=0, column=1, sticky="ns")

        self.outer = tk.Frame(self.canvas, bg=APP_BACKGROUND, padx=42, pady=34)
        outer = self.outer
        self.canvas_window = self.canvas.create_window(
            (0, 0),
            window=outer,
            anchor="nw",
        )
        outer.bind("<Configure>", self._update_scroll_region)
        self.canvas.bind("<Configure>", self._fit_content_width)
        self.canvas.bind_all("<MouseWheel>", self._scroll_page)

        header = tk.Frame(outer, bg=APP_BACKGROUND)
        header.pack(fill="x", pady=(0, 26))

        logo_shell = tk.Frame(header, bg=ACCENT_SOFT, padx=10, pady=10)
        logo_shell.pack(side="left", padx=(0, 16))
        logo = tk.Label(
            logo_shell,
            text="S",
            bg=ACCENT,
            fg="#FFFFFF",
            font=("Segoe UI Semibold", 17),
            width=2,
        )
        logo.pack()

        title_box = tk.Frame(header, bg=APP_BACKGROUND)
        title_box.pack(side="left")
        tk.Label(
            title_box,
            text="StarRez Contact Runner",
            bg=APP_BACKGROUND,
            fg=PRIMARY_TEXT,
            font=("Segoe UI Semibold", 24),
        ).pack(anchor="w")
        tk.Label(
            title_box,
            text="Add contacts to StarRez from an Excel file",
            bg=APP_BACKGROUND,
            fg=SECONDARY_TEXT,
            font=("Segoe UI", 10),
        ).pack(anchor="w")

        self.banner_labels = []
        for text, background, foreground in (
            ("Before you start: log in to StarRez and open the Main page.",
             ACCENT_SOFT, ACCENT),
            ("Close your selected Excel file before starting. Keep it closed while the script runs.",
             WARNING_SOFT, WARNING),
            ("Once the script starts, do not touch the mouse or keyboard until the entire list finishes.",
             WARNING_SOFT, WARNING),
        ):
            banner = tk.Label(
                outer, text=text, bg=background, fg=foreground,
                font=("Segoe UI", 17, "bold"), anchor="w", justify="left",
                padx=18, pady=16, wraplength=720,
            )
            banner.pack(fill="x", pady=(0, 14))
            self.banner_labels.append(banner)

        file_card = self._card(outer)
        file_card.pack(fill="x", pady=(0, 14))
        self._section_label(file_card, "1. Choose your Excel file").pack(anchor="w")
        tk.Label(
            file_card,
            text="Use column headings First, Last, and Email in the first row.",
            bg=CARD_BACKGROUND,
            fg=SECONDARY_TEXT,
            font=("Segoe UI", 10),
        ).pack(anchor="w", pady=(4, 12))

        self.example_button = self._button(
            file_card, "Click here to view an example Excel sheet",
            self._open_excel_example, subtle=True,
        )
        self.example_button.pack(anchor="w")

        file_row = tk.Frame(file_card, bg=CARD_BACKGROUND)
        file_row.pack(fill="x", pady=(10, 0))
        self.file_entry = tk.Entry(
            file_row,
            textvariable=self.excel_path,
            bg=FIELD_BACKGROUND,
            fg=PRIMARY_TEXT,
            insertbackground=PRIMARY_TEXT,
            disabledbackground=FIELD_BACKGROUND,
            disabledforeground=SECONDARY_TEXT,
            relief="solid",
            bd=1,
            highlightthickness=0,
            font=("Segoe UI", 10),
        )
        self.file_entry.pack(side="left", fill="x", expand=True, ipady=10)
        self.browse_button = self._button(
            file_row, "Choose file", self._choose_excel, subtle=True
        )
        self.browse_button.pack(side="left", padx=(10, 0), ipadx=8, ipady=5)

        status_card = self._card(outer)
        status_card.pack(fill="both", expand=True, pady=(0, 14))

        status_top = tk.Frame(status_card, bg=CARD_BACKGROUND)
        status_top.pack(fill="x")
        self._section_label(status_top, "Current contact").pack(side="left")
        self.status_badge = tk.Label(
            status_top,
            textvariable=self.status_text,
            bg=ACCENT_SOFT,
            fg=ACCENT,
            font=("Segoe UI Semibold", 9),
            padx=12,
            pady=5,
        )
        self.status_badge.pack(side="right")

        tk.Label(
            status_card,
            textvariable=self.current_contact,
            bg=CARD_BACKGROUND,
            fg=PRIMARY_TEXT,
            font=("Segoe UI Semibold", 22),
            anchor="w",
        ).pack(fill="x", pady=(22, 5))

        self.detail_label = tk.Label(
            status_card,
            text="Choose your file, then click Start contacts below.",
            bg=CARD_BACKGROUND,
            fg=SECONDARY_TEXT,
            font=("Segoe UI", 10),
            anchor="w",
            justify="left",
        )
        self.detail_label.pack(fill="x")

        progress_row = tk.Frame(status_card, bg=CARD_BACKGROUND)
        progress_row.pack(fill="x", pady=(24, 0))
        self.progress = ttk.Progressbar(
            progress_row,
            style="Runner.Horizontal.TProgressbar",
            mode="determinate",
            maximum=1,
            value=0,
        )
        self.progress.pack(side="left", fill="x", expand=True)
        tk.Label(
            progress_row,
            textvariable=self.progress_text,
            bg=CARD_BACKGROUND,
            fg=SECONDARY_TEXT,
            font=("Segoe UI Semibold", 9),
            width=8,
        ).pack(side="right", padx=(14, 0))

        hint = tk.Frame(status_card, bg="#F7F9FC", padx=16, pady=14)
        hint.pack(fill="x", pady=(24, 0))
        tk.Label(
            hint,
            text="2. Start your contacts",
            bg="#F7F9FC",
            fg=ACCENT,
            font=("Segoe UI Semibold", 9),
        ).pack(anchor="w")
        self.hint_text = tk.Label(
            hint,
            text="Close your Excel file, then click Start contacts. "
                 "You have 5 seconds to switch to the Main page in StarRez. "
                 "The app will check that StarRez is ready on your laptop or connected monitor. "
                 "All contacts will run automatically. Leave the mouse and keyboard alone until the list finishes.\n"
                 "To stop in an emergency, move the mouse to the top-left corner of the screen.",
            bg="#F7F9FC",
            fg=SECONDARY_TEXT,
            font=("Segoe UI", 9),
            anchor="w",
            justify="left",
        )
        self.hint_text.pack(fill="x", pady=(3, 0))

        controls = tk.Frame(outer, bg=APP_BACKGROUND)
        controls.pack(fill="x")
        self.saved_note = tk.Label(
            controls,
            text="We'll remember your Excel file for next time.",
            bg=APP_BACKGROUND,
            fg=SECONDARY_TEXT,
            font=("Segoe UI", 9),
            anchor="w",
            justify="left",
        )
        self.saved_note.pack(fill="x", pady=(0, 10))
        self.start_button = self._button(
            controls, "Start contacts", self._start
        )
        self.start_button.pack(fill="x", ipady=7)

    def _card(self, parent):
        return tk.Frame(
            parent,
            bg=CARD_BACKGROUND,
            padx=24,
            pady=22,
            highlightbackground=BORDER,
            highlightcolor=BORDER,
            highlightthickness=1,
            bd=0,
        )

    def _section_label(self, parent, text):
        return tk.Label(
            parent,
            text=text,
            bg=CARD_BACKGROUND,
            fg=PRIMARY_TEXT,
            font=("Segoe UI Semibold", 12),
        )

    def _button(self, parent, text, command, subtle=False):
        background = "#FFFFFF" if subtle else ACCENT
        foreground = PRIMARY_TEXT if subtle else "#FFFFFF"
        active_background = ACCENT_SOFT if subtle else ACCENT_HOVER
        return tk.Button(
            parent,
            text=text,
            command=command,
            bg=background,
            fg=foreground,
            activebackground=active_background,
            activeforeground=foreground,
            disabledforeground=SECONDARY_TEXT,
            relief="solid" if subtle else "flat",
            bd=1 if subtle else 0,
            cursor="hand2",
            font=("Segoe UI Semibold", 10),
            padx=15,
            pady=8,
        )

    def _open_excel_example(self):
        if self.running or self.example_page is not None:
            return
        image_path = Path(__file__).with_name("excel.png")
        self.example_source = None
        try:
            with Image.open(image_path) as source:
                self.example_source = source.convert("RGBA")
        except OSError:
            pass

        self.main_page.pack_forget()
        self.example_page = tk.Frame(self, bg=APP_BACKGROUND, padx=24, pady=24)
        self.example_page.pack(fill="both", expand=True)
        self._button(
            self.example_page, "Back to main page", self._close_excel_example,
            subtle=True,
        ).pack(anchor="w", pady=(0, 16))
        tk.Label(
            self.example_page, text="Example Excel sheet", bg=APP_BACKGROUND,
            fg=PRIMARY_TEXT, font=("Segoe UI", 20, "bold"),
        ).pack(anchor="w", pady=(0, 8))
        tk.Label(
            self.example_page,
            text="Use First, Last, and Email as your column headings.",
            bg=APP_BACKGROUND, fg=SECONDARY_TEXT, font=("Segoe UI", 10),
            wraplength=440, justify="left",
        ).pack(anchor="w", pady=(0, 16))
        self.example_display = tk.Label(
            self.example_page, bg=CARD_BACKGROUND, fg=SECONDARY_TEXT,
            font=("Segoe UI", 12), wraplength=400,
        )
        self.example_display.pack(fill="both", expand=True)
        if self.example_source is None:
            self.example_display.config(
                text="The example image isn't available yet.\n\n"
                     "Add an image named excel.png to the app's folder.",
            )
        else:
            self.example_display.bind("<Configure>", self._resize_excel_example)

    def _resize_excel_example(self, event):
        preview = self.example_source.copy()
        preview.thumbnail((max(1, event.width - 24), max(1, event.height - 24)),
                          Image.Resampling.LANCZOS)
        self.example_photo = ImageTk.PhotoImage(preview, master=self)
        self.example_display.config(image=self.example_photo)

    def _close_excel_example(self):
        self.example_page.destroy()
        self.example_page = None
        self.example_source = None
        self.example_photo = None
        self.main_page.pack(fill="both", expand=True)

    def _choose_excel(self):
        selected = filedialog.askopenfilename(
            title="Choose your Excel contact file",
            initialdir=str(Path(self.excel_path.get()).parent),
            filetypes=[("Excel workbooks", "*.xlsx"), ("All files", "*.*")],
        )
        if selected:
            self.excel_path.set(selected)
            self._save_excel_path()
            self.status_text.set("Ready")
            self.detail_label.config(
                text="File selected. Log in to StarRez and open Main, then click Start contacts.")

    def _on_resize(self, event):
        if event.widget is not self:
            return
        side_padding = 24 if event.width < 780 else 42
        self.outer.config(padx=side_padding)
        text_width = max(280, event.width - (side_padding * 2) - 80)
        self.detail_label.config(wraplength=text_width)
        self.hint_text.config(wraplength=text_width)
        self.saved_note.config(wraplength=text_width)
        for banner in self.banner_labels:
            banner.config(wraplength=text_width)

    def _update_scroll_region(self, _event=None):
        self.canvas.configure(scrollregion=self.canvas.bbox("all"))

    def _fit_content_width(self, event):
        self.canvas.itemconfigure(self.canvas_window, width=event.width)

    def _scroll_page(self, event):
        if self.example_page is not None:
            return
        if self.canvas.winfo_height() < self.outer.winfo_reqheight():
            self.canvas.yview_scroll(int(-event.delta / 120), "units")

    def _start(self):
        if self.running:
            return
        workbook_path = Path(self.excel_path.get().strip())
        if not workbook_path.is_file():
            messagebox.showerror(
                "File not found", "Click Choose file and select your Excel file.")
            return

        try:
            workbook = load_workbook(
                workbook_path, read_only=True, data_only=True)
            try:
                get_column_numbers(workbook.active)
            finally:
                workbook.close()
        except (OSError, ValueError) as error:
            messagebox.showerror("Cannot open this Excel file", str(error))
            return

        self.running = True
        self.cancel_run.clear()
        self.website_retry.clear()
        self._save_excel_path()
        self.start_button.config(state="disabled")
        self.browse_button.config(state="disabled")
        self.example_button.config(state="disabled")
        self.file_entry.config(state="disabled")
        self.status_text.set("Loading")
        self.status_badge.config(fg=ACCENT, bg=ACCENT_SOFT)

        worker = threading.Thread(
            target=self._run_contacts,
            args=(workbook_path,),
            daemon=True,
        )
        worker.start()

    def _wait_for_starrez(self, countdown=False):
        """Pause the worker, keeping Tk responsive, until readiness is verified."""
        while not self.cancel_run.is_set():
            if countdown:
                for seconds in range(5, 0, -1):
                    self.event_queue.put(("countdown", seconds))
                    if self.cancel_run.wait(1):
                        return None
            self.event_queue.put(("checking_website",))
            region, reason = check_starrez_ready()
            if self.cancel_run.is_set():
                return None
            if region is not None:
                return region
            self.website_retry.clear()
            self.event_queue.put(("website_required", reason))
            self.website_retry.wait()
            countdown = True
        return None

    def _show_website_prompt(self, reason):
        self.status_text.set("Paused: open StarRez")
        self.status_badge.config(fg=WARNING, bg=WARNING_SOFT)
        self.detail_label.config(text=reason)
        self.deiconify()
        self.lift()
        prompt = tk.Toplevel(self)
        self.website_prompt = prompt
        prompt.title("Open StarRez to continue")
        prompt.configure(bg=WARNING_SOFT, padx=24, pady=24)
        prompt.resizable(False, False)
        prompt.transient(self)
        tk.Label(
            prompt, text="Please open StarRez before continuing",
            bg=WARNING_SOFT, fg=WARNING, font=("Segoe UI", 16, "bold"),
            wraplength=440, justify="left",
        ).pack(anchor="w", pady=(0, 12))
        tk.Label(
            prompt, text=reason,
            bg=WARNING_SOFT, fg=PRIMARY_TEXT, font=("Segoe UI", 11),
            wraplength=440, justify="left",
        ).pack(anchor="w", pady=(0, 12))
        copy_note = tk.Label(
            prompt, text="Click the link to copy it.",
            bg=WARNING_SOFT, fg=SECONDARY_TEXT, font=("Segoe UI", 10),
            wraplength=440, justify="left",
        )
        tk.Button(
            prompt, text=STARREZ_URL,
            command=lambda: self._copy_website_link(copy_note),
            bg=WARNING_SOFT, fg=ACCENT,
            activebackground=WARNING_SOFT, activeforeground=ACCENT_HOVER,
            font=("Segoe UI", 11, "underline"), cursor="hand2",
            relief="flat", bd=0, padx=0, pady=4,
            wraplength=440, justify="left", anchor="w",
        ).pack(anchor="w", fill="x")
        copy_note.pack(anchor="w", pady=(0, 16))
        tk.Label(
            prompt, text="You can use your laptop screen or either monitor.\n\n"
                         "When ready, click the button below. You will have 5 seconds "
                         "to switch back to StarRez before we check again.",
            bg=WARNING_SOFT, fg=PRIMARY_TEXT, font=("Segoe UI", 11),
            wraplength=440, justify="left",
        ).pack(anchor="w", pady=(0, 20))
        self._button(prompt, "I am on the website now", self._retry_website).pack(fill="x")
        self._button(prompt, "Cancel run", self._cancel_website_wait, subtle=True).pack(
            fill="x", pady=(10, 0))
        prompt.protocol("WM_DELETE_WINDOW", self._cancel_website_wait)

    def _copy_website_link(self, feedback):
        try:
            self.clipboard_clear()
            self.clipboard_append(STARREZ_URL)
        except tk.TclError:
            feedback.config(text="Couldn't copy the link. Please try clicking it again.", fg=ERROR)
            return
        feedback.config(text="Copied! Paste it into your browser's address bar with Ctrl+V.", fg=SUCCESS)

    def _dismiss_website_prompt(self):
        if self.website_prompt is not None:
            self.website_prompt.destroy()
            self.website_prompt = None

    def _retry_website(self):
        self._dismiss_website_prompt()
        self.status_text.set("Getting ready")
        self.website_retry.set()

    def _cancel_website_wait(self):
        self._dismiss_website_prompt()
        self.cancel_run.set()
        self.website_retry.set()

    def _close_app(self):
        self.cancel_run.set()
        self.website_retry.set()
        self.destroy()

    def _run_contacts(self, workbook_path):
        workbook = None
        try:
            workbook = load_workbook(
                workbook_path, read_only=True, data_only=True)
            sheet = workbook.active
            columns = get_column_numbers(sheet)
            contacts = []

            row_number = 2
            while row_number <= sheet.max_row:
                first_name = str(
                    sheet.cell(row=row_number,
                               column=columns["First"]).value or ""
                ).strip()
                last_name = str(
                    sheet.cell(row=row_number,
                               column=columns["Last"]).value or ""
                ).strip()
                email = str(
                    sheet.cell(row=row_number,
                               column=columns["Email"]).value or ""
                ).strip()

                if first_name and last_name and email:
                    contacts.append((row_number, first_name, last_name, email))
                row_number += 1

            total = len(contacts)
            self.event_queue.put(("total", total))

            for position, (row_number, first_name, last_name, email) in enumerate(
                contacts, start=1
            ):
                region = self._wait_for_starrez(countdown=(position == 1))
                if region is None:
                    self.event_queue.put(("stopped", "Run cancelled. No more contacts will be added."))
                    return
                self.event_queue.put(
                    ("running", position, total, row_number,
                     first_name, last_name, email)
                )

                try:
                    process_contact(first_name, last_name, email, screen_region=region)
                    self.event_queue.put(("finished", position, total))
                except pyautogui.ImageNotFoundException as error:
                    error_text = str(
                        error) or "A required screen image was not found."
                    self.event_queue.put(
                        ("failed", position, total, error_text))
                    return

            self.event_queue.put(("complete", total))
        except pyautogui.FailSafeException:
            self.event_queue.put(("stopped", "Emergency stop triggered."))
        except Exception as error:
            self.event_queue.put(("fatal", f"{type(error).__name__}: {error}"))
        finally:
            if workbook is not None:
                workbook.close()

    def _process_events(self):
        try:
            while True:
                event = self.event_queue.get_nowait()
                kind = event[0]

                if kind == "total":
                    total = event[1]
                    self.progress.config(maximum=max(total, 1), value=0)
                    self.progress_text.set(f"0 / {total}")
                elif kind == "countdown":
                    self.status_text.set(f"Starting in {event[1]} seconds")
                    self.status_badge.config(fg=WARNING, bg=WARNING_SOFT)
                    self.detail_label.config(
                        text="Switch to the Main page in StarRez now. Then leave the mouse and keyboard alone.")
                elif kind == "checking_website":
                    self.status_text.set("Checking StarRez")
                    self.detail_label.config(text="Keep StarRez in front while we check that the Main page is ready.")
                elif kind == "website_required":
                    self._show_website_prompt(event[1])
                elif kind == "running":
                    _, position, total, row, first, last, email = event
                    self.current_contact.set(f"{first} {last}")
                    self.status_text.set("Running")
                    self.status_badge.config(fg=ACCENT, bg=ACCENT_SOFT)
                    self.detail_label.config(
                        text="Adding this contact to StarRez. Do not touch the mouse or keyboard.")
                    self.progress_text.set(f"{position - 1} / {total}")
                elif kind == "finished":
                    _, position, total = event
                    self.progress.config(value=position)
                    self.progress_text.set(f"{position} / {total}")
                    self.status_text.set("Contact finished")
                    self.status_badge.config(fg=SUCCESS, bg=SUCCESS_SOFT)
                elif kind == "failed":
                    _, position, total, error_text = event
                    self.progress.config(value=position)
                    self.progress_text.set(f"{position} / {total}")
                    self.status_text.set("Stopped: contact could not finish")
                    self.status_badge.config(fg=ERROR, bg=ERROR_SOFT)
                    self.detail_label.config(
                        text=f"{error_text}\nCheck this contact in StarRez before starting again. "
                             "Remove completed contacts from your Excel list to avoid repeating them.")
                    self._unlock_controls()
                elif kind == "complete":
                    total = event[1]
                    self.current_contact.set("Contact list finished")
                    self.status_text.set("Finished")
                    self.status_badge.config(fg=SUCCESS, bg=SUCCESS_SOFT)
                    self.detail_label.config(
                        text=f"All {total} valid contacts were checked.")
                    self._unlock_controls()
                elif kind in ("stopped", "fatal"):
                    self.status_text.set(
                        "Stopped" if kind == "stopped" else "Error")
                    self.status_badge.config(fg=ERROR, bg=ERROR_SOFT)
                    self.detail_label.config(text=event[1])
                    self._unlock_controls()
        except queue.Empty:
            pass

        self.after(100, self._process_events)

    def _unlock_controls(self):
        self.running = False
        self.start_button.config(state="normal")
        self.browse_button.config(state="normal")
        self.example_button.config(state="normal")
        self.file_entry.config(state="normal")


if __name__ == "__main__":
    app = AutomationLauncher()
    app.mainloop()
